import os
import io
import subprocess
import zipfile
from datetime import datetime
from user_agents import parse
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from dotenv import load_dotenv

from docx import Document
from pymongo import MongoClient
from pymongo.errors import ServerSelectionTimeoutError
import certifi
from docx.enum.style import WD_STYLE_TYPE
import uuid



load_dotenv()

app = Flask(__name__)
CORS(app, origins=["http://localhost:4200", "https://ryanbeevers.github.io*", "http://raspberrypi.local:4200"])

LLAMA_BIN = "/home/ryan2914/llama.cpp/build/bin/llama-run"
MODEL_PATH = "file:///home/ryan2914/llama.cpp/models/tinyllama-1.1b-chat-v1.0.Q4_K_M.gguf"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")
OUTPUT_DIR = os.path.join(BASE_DIR, "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)

try:
    client = MongoClient(
        os.environ.get("MONGO_URI"),
        tls=True,
        tlsCAFile=certifi.where()
    )
    client.server_info()
except ServerSelectionTimeoutError as err:
    print("Failed to connect to MongoDB:", err)
    exit(1)

db = client["resume_db"]
visits_col = db["visits"]
progress_col = db["progress"]

@app.route('/track-visit', methods=['POST'])
def track_visit():
    data = request.json

    forwarded_for = request.headers.get('X-Forwarded-For', request.remote_addr)
    ip_address = forwarded_for.split(',')[0].strip()
    ua_string = request.headers.get('User-Agent', '')
    user_agent = parse(ua_string)

    visit = {
        "user_id": data.get("user_id"),
        "username": data.get("username") or "anonymous",
        "ip_address": ip_address,
        "user_agent": request.headers.get('User-Agent'),
        "accept_language": request.headers.get('Accept-Language'),
        "referer": request.headers.get('Referer'),
        "platform": user_agent.os.family,
        "browser": user_agent.browser.family,
        "version": user_agent.browser.version_string,
        "mobile": user_agent.is_mobile,
        "timestamp": datetime.utcnow()
    }

    visits_col.insert_one(visit)
    return jsonify({"status": "visit logged"})

all_star_ids = {"star1", "star2", "star3", "star4", "star5"}

@app.route('/progress/star', methods=['POST'])
def add_star():
    data = request.json
    user_id = data.get("user_id")
    star_id = data.get("star_id")

    if not user_id or not star_id:
        return jsonify({"error": "Missing user_id or star_id"}), 400

    progress = progress_col.find_one({"user_id": user_id}) or {"stars_found": []}
    updated_stars = set(progress.get("stars_found", []))
    updated_stars.add(star_id)

    update_data = {
        "$addToSet": {"stars_found": star_id},
        "$set": {"last_updated": datetime.utcnow()}
    }

    if updated_stars == all_star_ids and not progress.get("completed"):
        update_data["$set"].update({
            "completed": True,
            "completed_at": datetime.utcnow()
        })

    result = progress_col.update_one(
        {"user_id": user_id},
        update_data,
        upsert=True
    )

    return jsonify({
        "status": "star added",
        "modified_count": result.modified_count,
        "completed": updated_stars == all_star_ids
    })


@app.route('/progress/<user_id>', methods=['GET'])
def get_progress(user_id):
    progress = progress_col.find_one({"user_id": user_id}, {"_id": 0})
    if not progress:
        return jsonify({"stars_found": []})
    return jsonify(progress)

@app.route('/progress/status/<user_id>', methods=['GET'])
def progress_status(user_id):
    progress = progress_col.find_one({"user_id": user_id}, {"_id": 0, "completed": 1})
    return jsonify({"completed": progress.get("completed", False)})

@app.route('/progress/complete', methods=['POST'])
def mark_complete():
    data = request.json
    user_id = data.get("user_id")

    if not user_id:
        return jsonify({"error": "Missing user_id"}), 400

    result = progress_col.update_one(
        {"user_id": user_id},
        {
            "$set": {
                "completed": True,
                "completed_at": datetime.utcnow()
            }
        },
        upsert=True
    )
    return jsonify({"status": "completed", "modified_count": result.modified_count})
@app.route('/user-summary', methods=['GET'])
def get_user_summary():
    pipeline = [
        {"$sort": {"timestamp": -1}},  # Ensures latest values appear first
        {
            "$group": {
                "_id": "$user_id",
                "visit_count": {"$sum": 1},
                "last_visit": {"$first": "$timestamp"},
                "ip_address": {"$first": "$ip_address"},
                "username": {"$first": "$username"},
                "platform": {"$first": "$platform"},
                "browser": {"$first": "$browser"},
                "version": {"$first": "$version"},
                "mobile": {"$first": "$mobile"},
            }
        },
        {"$sort": {"last_visit": -1}}
    ]

    summary = list(visits_col.aggregate(pipeline))
    for item in summary:
        item["user_id"] = item.pop("_id")
        if isinstance(item["last_visit"], datetime):
            item["last_visit"] = item["last_visit"].isoformat()

    return jsonify({"summary": summary})

@app.route('/visits/<user_id>', methods=['GET'])
def get_visits_by_user(user_id):
    visits = list(visits_col.find({"user_id": user_id}).sort("timestamp", -1))
    for visit in visits:
        visit["_id"] = str(visit["_id"])
        if isinstance(visit["timestamp"], datetime):
            visit["timestamp"] = visit["timestamp"].isoformat()
    return jsonify({"visits": visits})

@app.route('/generate', methods=['POST'])
def generate_resume_text():
    data = request.json
    prompt = data.get("prompt", "")
    if not prompt:
        return jsonify({"error": "Missing prompt"}), 400

    try:
        result = subprocess.run(
            [LLAMA_BIN, MODEL_PATH, prompt],
            capture_output=True, text=True, timeout=300
        )
        return jsonify({"response": result.stdout.strip()})
    except subprocess.TimeoutExpired:
        return jsonify({"error": "Model timed out"}), 504
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ------------------------------
# Utility functions
# ------------------------------
def replace_placeholder(doc, placeholder, replacement):
    replaced = False

    # Paragraphs
    for paragraph in doc.paragraphs:
        if placeholder in paragraph.text:
            paragraph.text = paragraph.text.replace(placeholder, replacement)
            replaced = True

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if placeholder in cell.text:
                    cell.text = cell.text.replace(placeholder, replacement)
                    replaced = True

    if not replaced:
        print(f"Warning: Placeholder not found: {placeholder}")

    return doc


def add_bullets(doc, placeholder, bullets):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    style_name = "List Bullet"
    # Check if style exists in the document
    if style_name not in [s.name for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH]:
        style_name = "Normal"

    for paragraph in doc.paragraphs:
        if placeholder in paragraph.text:
            paragraph.text = ""
            parent = paragraph._p.getparent()
            idx = parent.index(paragraph._p)
            for bullet in bullets:
                new_p = OxmlElement("w:p")
                parent.insert(idx + 1, new_p)
                para = Paragraph(new_p, paragraph._parent)
                para.style = style_name
                para.add_run(bullet)
                idx += 1
            break
    return doc


def validate_docx(path):
    with ZipFile(path, "r") as z:
        bad = z.testzip()
        if bad:
            raise ValueError(f"Corrupt DOCX zip entry: {bad}")


# ------------------------------
# Resume / Cover Builder
# ------------------------------
def build_resume(json_data, template_path, output_dir):
    doc = Document(template_path)

    # Header
    doc = replace_placeholder(doc, "{{NAME}}", json_data["header"]["name"])
    doc = replace_placeholder(doc, "{{TITLE}}", json_data["header"]["title"])
    doc = replace_placeholder(doc, "{{TAGLINE}}", json_data["header"]["tagline"])
    doc = replace_placeholder(doc, "{{LOCATION_NOTE}}", json_data["header"]["location_note"])
    doc = replace_placeholder(doc, "{{SUMMARY}}", json_data["summary"])

    # Experience
    for idx, exp in enumerate(json_data["experience"]):
        placeholder = f"{{{{EXP_{idx + 1}}}}}"
        doc = add_bullets(doc, placeholder, exp["bullets"])

    # Skills
    skills = json_data.get("skills", {})
    doc = replace_placeholder(doc, "{{FRONTEND_SKILLS}}", ", ".join(skills.get("frontend", [])))
    doc = replace_placeholder(doc, "{{BACKEND_SKILLS}}", ", ".join(skills.get("backend", [])))
    doc = replace_placeholder(doc, "{{CLOUD_DEVOPS_SKILLS}}", ", ".join(skills.get("cloud_devops", [])))
    doc = replace_placeholder(doc, "{{SECURITY_AUTH_SKILLS}}", ", ".join(skills.get("security_auth", [])))
    doc = replace_placeholder(doc, "{{DATABASE_SKILLS}}", ", ".join(skills.get("databases", [])))
    doc = replace_placeholder(doc, "{{TOOLS_SKILLS}}", ", ".join(skills.get("tools_collaboration", [])))

    filename = f"resume_{uuid.uuid4()}.docx"
    output_path = os.path.join(output_dir, filename)
    doc.save(output_path)
    validate_docx(output_path)
    return filename


def build_cover_letter(json_data, template_path, output_dir):
    doc = Document(template_path)

    doc = replace_placeholder(doc, "{{COMPANY}}", json_data["recipient"]["company"])
    doc = replace_placeholder(doc, "{{ROLE}}", json_data["recipient"]["role"])
    doc = replace_placeholder(doc, "{{OPENING_PARAGRAPH}}", json_data["opening_paragraph"])

    for i, para in enumerate(json_data.get("body_paragraphs", [])):
        placeholder = f"{{{{BODY_PARAGRAPH_{i + 1}}}}}"
        doc = replace_placeholder(doc, placeholder, para)

    doc = replace_placeholder(doc, "{{CLOSING_PARAGRAPH}}", json_data["closing_paragraph"])
    doc = replace_placeholder(doc, "{{NAME}}", json_data["signature"]["name"])

    filename = f"cover_letter_{uuid.uuid4()}.docx"
    output_path = os.path.join(output_dir, filename)
    doc.save(output_path)
    validate_docx(output_path)
    return filename


# ------------------------------
# Flask Endpoints
# ------------------------------
@app.route("/generate_docs", methods=["POST"])
def generate_docs():
    data = request.json

    resume_template = os.path.join(TEMPLATE_DIR, "resume_template.docx")
    cover_template = os.path.join(TEMPLATE_DIR, "cover_letter_template.docx")

    # Build docs
    resume_file = build_resume(data["resume"], resume_template, OUTPUT_DIR)
    cover_file = build_cover_letter(data["cover_letter"], cover_template, OUTPUT_DIR)

    # ZIP in memory
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zipf:
        zipf.write(os.path.join(OUTPUT_DIR, resume_file), arcname="Resume.docx")
        zipf.write(os.path.join(OUTPUT_DIR, cover_file), arcname="Cover Letter.docx")

    zip_buffer.seek(0)

    # Filename: 01-12-2026 - Company Name.zip
    company = data["cover_letter"]["recipient"]["company"]
    safe_company = "".join(c for c in company if c.isalnum() or c in " -_").strip()
    date_str = datetime.now().strftime("%m-%d-%Y")
    zip_filename = f"{date_str} - {safe_company}.zip"

    return send_file(
        zip_buffer,
        mimetype="application/zip",
        as_attachment=True,
        download_name=zip_filename,
    )


@app.route("/download/<filename>", methods=["GET"])
def download_file(filename):
    return send_file(os.path.join(OUTPUT_DIR, filename), as_attachment=True)



if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port, debug=True)

